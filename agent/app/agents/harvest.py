"""Find the questions in a buyer's workbook, and where each answer goes.

Buyer sheets come in a handful of shapes, and a single "answer is one column to the right" rule fails
on most of them:

* Ref | Question | Answer | Notes: the answer sits right of the question
* Item | Unit | Answer: a unit column between question and answer
* a header row (回答, Response, Giá trị) that names the answer column, merged question cells that
  span several columns, so the answer is right of the *merge*, not right of its first cell
* a question that spans the page, with the answer in the row below it
* energy tables where the "question" is just a short row label ("Diesel", "電力") and a
  header says which column takes the quantity
* an answer cell that is a dropdown: only one of its options is a valid answer

The harvest decides all of that per row, by looking at the sheet, not by asking a model. The model
only maps wording to fields afterwards. The row stays a deterministic fact: where the answer is
written is never the model's choice.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from typing import Any

from openpyxl.utils import get_column_letter, range_boundaries

ANSWER_KEYS = (
    "answer", "response", "reply", "your input", "input", "value", "actual", "result", "amount", "quantity",
    "回答", "記入", "入力", "実績", "数値", "使用量", "消費量", "値", "trả lời", "phản hồi", "giá trị", "số liệu",
)
UNIT_KEYS = ("unit", "uom", "単位", "đơn vị", "don vi")
NOTE_KEYS = ("note", "remark", "comment", "evidence", "source", "備考", "コメント", "根拠", "出典", "ghi chú", "nguồn", "bằng chứng")

_EMPTY_MARKS = {"", "-", "–", "—", "－", "ー", "n/a", "###", "tbd"}


def is_empty(v: Any) -> bool:
    if v is None:
        return True
    s = str(v).strip().lower()
    return s in _EMPTY_MARKS or bool(re.fullmatch(r"[_.…\s]+", s))


@dataclass
class QuestionCell:
    id: str
    sheet: str
    row: int
    col: int
    text: str
    ref: str | None
    # Where the answer is written. Set for every harvested question.
    answer_row: int = 0
    answer_col: int = 0
    layout: str = "right"           # right | below
    unit: str | None = None         # the unit the sheet asks for, as printed (kWh, t-CO2, L)
    section: str | None = None
    # The numbered question this row sits under, for a table question laid out as one row per column.
    parent: str | None = None
    options: list[str] = field(default_factory=list)   # dropdown choices at the answer cell, if any


def _looks_like_ref(v: Any) -> bool:
    s = str(v).strip()
    return 0 < len(s) <= 10 and any(ch.isdigit() for ch in s) and " " not in s


def _styled_as_heading(cell) -> bool:
    """A lone text counts as a section heading only when it is bold or filled, as headings are. A lone
    label inside a table (one column of a numbered question) is a row, not a heading."""
    try:
        return bool(cell.font and cell.font.b) or bool(cell.fill and cell.fill.fill_type)
    except AttributeError:
        return False


def _has(text: str, keys: tuple[str, ...]) -> bool:
    t = text.lower()
    return any(k in t for k in keys)


@dataclass
class _Header:
    row: int = 0
    answer_cols: list[int] = field(default_factory=list)
    unit_cols: list[int] = field(default_factory=list)
    note_cols: list[int] = field(default_factory=list)


def _merged(ws) -> dict[tuple[int, int], tuple[int, int, int, int]]:
    """(row, col) of every cell inside a merge -> (min_row, min_col, max_row, max_col)."""
    out: dict[tuple[int, int], tuple[int, int, int, int]] = {}
    for rng in ws.merged_cells.ranges:
        min_col, min_row, max_col, max_row = range_boundaries(str(rng))
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                out[(r, c)] = (min_row, min_col, max_row, max_col)
    return out


def _find_header(ws, merged) -> _Header:
    """The first row near the top that has an answer-ish heading beside other short headings."""
    for r in range(1, min(ws.max_row, 30) + 1):
        texts: list[tuple[int, str]] = []
        for c in range(1, min(ws.max_column, 30) + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip() and ((r, c) not in merged or merged[(r, c)][:2] == (r, c)):
                texts.append((c, v.strip()))
        if len(texts) < 2 or any(len(t) > 40 for _, t in texts):
            continue
        ans = [c for c, t in texts if _has(t, ANSWER_KEYS) and not _has(t, UNIT_KEYS)]
        if not ans:
            continue
        return _Header(
            r, ans,
            [c for c, t in texts if _has(t, UNIT_KEYS)],
            # a note column sits right of the answer; "Energy source" in the first column is the question column
            [c for c, t in texts if _has(t, NOTE_KEYS) and c not in ans and c > min(ans)],
        )
    return _Header()


def _text_unit(text: str) -> str | None:
    """The unit printed in a question, '(kWh)', '（MWh）', 'in metric tons CO2e', or None."""
    m = re.search(r"[(（\[]\s*([^()（）\[\]]{1,24}?)\s*[)）\]]\s*[?？:：]?\s*$", text)
    if m and re.search(r"(kwh|mwh|gj|mj|tj|\bkg\b|\bt\b|\bl\b|litre|liter|m3|m³|co2|トン|リットル|㎥)", m.group(1), re.IGNORECASE):
        return m.group(1).strip()
    m = re.search(r"\bin\s+(metric tons|tonnes|tons|kg|mwh|kwh|gj|litres|liters)\b(\s+co2e?)?", text, re.IGNORECASE)
    if m:
        return (m.group(1) + (m.group(2) or "")).strip()
    return None


def dropdown_options(ws, wb, coord: str) -> list[str]:
    """The valid choices at `coord` when it carries a list validation, else []."""
    dvs = getattr(ws, "data_validations", None)
    if not dvs:
        return []
    for dv in dvs.dataValidation:
        if dv.type != "list" or not dv.formula1 or coord not in dv.sqref:
            continue
        f = dv.formula1.strip()
        if f.startswith('"') and f.endswith('"'):
            return [x.strip() for x in f[1:-1].split(",") if x.strip()]
        ref = f.lstrip("=")
        if ref in wb.defined_names:
            ref = wb.defined_names[ref].attr_text
        sheet = ws
        if "!" in ref:
            name, ref = ref.rsplit("!", 1)
            name = name.strip("'")
            if name not in wb.sheetnames:
                return []
            sheet = wb[name]
        try:
            min_col, min_row, max_col, max_row = range_boundaries(ref.replace("$", ""))
        except ValueError:
            return []
        vals = [sheet.cell(r, c).value for r in range(min_row, max_row + 1) for c in range(min_col, max_col + 1)]
        return [str(v).strip() for v in vals if v not in (None, "")]
    return []


class _Sheet:
    def __init__(self, ws, merged) -> None:
        self.ws, self.merged = ws, merged

    def anchor(self, r: int, c: int) -> tuple[int, int]:
        m = self.merged.get((r, c))
        return (m[0], m[1]) if m else (r, c)

    def value(self, r: int, c: int):
        ar, ac = self.anchor(r, c)
        return self.ws.cell(ar, ac).value

    def span(self, r: int, c: int) -> tuple[int, int]:
        m = self.merged.get((r, c))
        return (m[1], m[3]) if m else (c, c)


def harvest(wb) -> list[QuestionCell]:
    cells: list[QuestionCell] = []
    n = 0
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        merged = _merged(ws)
        hdr = _find_header(ws, merged)
        a_cols, u_cols, n_cols = set(hdr.answer_cols), set(hdr.unit_cols), set(hdr.note_cols)
        first_answer = min(a_cols) if a_cols else None
        section: str | None = None
        parent: str | None = None

        sh = _Sheet(ws, merged)
        anchor, cell_value, span = sh.anchor, sh.value, sh.span

        for row in ws.iter_rows():
            r = row[0].row
            texts = [c for c in row if isinstance(c.value, str) and c.value.strip()]
            if hdr.row and r <= hdr.row:
                continue
            # a lone short text on a row is a section heading, kept as context for the model
            if len(texts) == 1 and len(texts[0].value.strip()) <= 60 and _styled_as_heading(texts[0]) and is_empty(
                    ws.cell(r, texts[0].column + 1).value if texts[0].column < ws.max_column else None):
                section = texts[0].value.strip()
                parent = None
            for cell in texts:
                c = cell.column
                if (r, c) in merged and merged[(r, c)][:2] != (r, c):
                    continue                                    # covered by a merge: its anchor speaks for it
                if c in a_cols or c in u_cols or c in n_cols:
                    continue
                if first_answer is not None and c >= first_answer:
                    continue
                text = cell.value.strip()
                left = ws.cell(r, c - 1).value if c > 1 else None
                ref = str(left).strip() if left is not None and _looks_like_ref(left) else None
                c0, c1 = span(r, c)
                right_val = cell_value(r, c1 + 1)
                # a bare reference number ("20.4") with its question text beside it is the ref, not a question
                if _looks_like_ref(text) and isinstance(right_val, str) and len(right_val.strip()) > 6:
                    continue
                short_label_ok = bool(a_cols) and len(text) >= 2
                if not (ref or text.endswith(("?", "？")) or len(text) > 25 or short_label_ok):
                    continue
                if ref:
                    parent_here, parent = parent, text
                else:
                    parent_here = parent

                unit = None
                ans_col = 0
                layout = "right"
                # an empty merged block of exactly the question's width right under it is an answer box
                box = (r + 1, c0) in merged and merged[(r + 1, c0)][1::2] == (c0, c1) and merged[(r + 1, c0)][0] == r + 1
                if box and is_empty(ws.cell(r + 1, c0).value):
                    pass
                elif a_cols:
                    cand = [x for x in sorted(a_cols) if x > c1]
                    ans_col = anchor(r, cand[0])[1] if cand else 0
                    if ans_col and not is_empty(ws.cell(r, ans_col).value):
                        ans_col = 0                              # already answered: leave the buyer's own entry alone
                    for u in sorted(u_cols):
                        v = ws.cell(r, u).value
                        if isinstance(v, str) and v.strip():
                            unit = v.strip()
                            break
                else:
                    for x in range(c1 + 1, c1 + 5):
                        if (r, x) in merged and merged[(r, x)][:2] != (r, x):
                            continue
                        v = ws.cell(r, x).value
                        if is_empty(v):
                            ans_col = x
                            break
                        if isinstance(v, str) and len(v.strip()) <= 12 and unit is None:
                            unit = v.strip()                     # a short word between question and answer is the unit
                            continue
                        break                                    # prose or a number: not an answer slot
                if not ans_col:
                    # the question spans the page, or has nothing free beside it: the answer is the row below
                    br, bc = anchor(r + 1, c0)
                    wide = (c1 - c0 >= 2) or ((r + 1, c0) in merged)
                    if wide and is_empty(ws.cell(br, bc).value):
                        n += 1
                        cells.append(QuestionCell(
                            f"q{n}", ws.title, r, c, text, ref, br, bc, "below", unit or _text_unit(text), section,
                            None if ref else parent_here, dropdown_options(ws, wb, f"{get_column_letter(bc)}{br}"),
                        ))
                    continue
                n += 1
                coord = f"{get_column_letter(ans_col)}{r}"
                cells.append(QuestionCell(
                    f"q{n}", ws.title, r, c, text, ref, r, ans_col, layout, unit or _text_unit(text), section,
                    None if ref else parent_here, dropdown_options(ws, wb, coord),
                ))
    return cells

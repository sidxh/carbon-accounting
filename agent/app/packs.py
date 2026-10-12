"""Answer packs for questionnaires that are answered in a portal.

CDP, Higg FEM and EcoVadis are answered inside the buyer's or the publisher's platform, so there is
no spreadsheet to write into. For those we hand the supplier a pack: every question in the official
bank, in the publisher's numbering and wording, with the answer to type in, the unit, how the number was
computed, and which documents to attach as evidence.

Every answer comes from `facts.compute_fact`, the same place the Excel filler reads from, so a number
in a pack is the number in a filled sheet. Nothing is invented: a question we cannot answer from bills
is listed as "answer in the portal", and an answer that needs the supplier's confirmation is flagged.
"""

from __future__ import annotations

import io
from datetime import UTC, datetime
from pathlib import Path

from jinja2 import Environment, FileSystemLoader, select_autoescape
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field

from .agents.filler import choose_option
from .exporters.methodology import _font_path, html_to_pdf
from .facts import FactContext, compute_fact
from .models import InventoryLine
from .questionnaires import Bank, BankQuestion, Source, get_bank

HERE = Path(__file__).resolve().parent / "exporters"
env = Environment(loader=FileSystemLoader(str(HERE)), autoescape=select_autoescape(["html"]))

PACK_BANKS = ("cdp_sme", "higg_fem_energy", "ecovadis_carbon")

CALC_EVIDENCE = "Calculation lines (the “Calculation lines” sheet of the spreadsheet export)"
METHOD_EVIDENCE = "Methodology note (PDF, from the Questionnaire page)"

GREEN, YELLOW, GREY = "E3F1E4", "FFF2CC", "ECECEC"


class PackCell(BaseModel):
    label: str | None = None
    value: str | float | int | None = None
    unit: str | None = None
    source: str = ""
    flagged: bool = False
    reason: str | None = None
    portal: bool = False        # only the supplier can answer this one, in the portal
    options: list[str] = Field(default_factory=list)    # the portal's own answer list, when the publisher prints one


class PackItem(BaseModel):
    ref: str
    ref_id: str | None = None
    text: str
    wording: str
    format: str
    status: str                 # answered | needs_you | portal
    cells: list[PackCell]
    evidence: list[str] = Field(default_factory=list)
    note: str | None = None


class Pack(BaseModel):
    bank_id: str
    name: str
    publisher: str
    version: str
    wording: str
    sources: list[Source]
    notes: list[str]
    company: str
    period: str
    generated_at: str
    items: list[PackItem]
    evidence_files: list[str]

    @property
    def counts(self) -> dict[str, int]:
        out = {"answered": 0, "needs_you": 0, "portal": 0}
        for i in self.items:
            out[i.status] += 1
        return out


# --- evidence -----------------------------------------------------------------------------------------

def _relevant(fld: str | None, energy_type: str | None, lines: list[InventoryLine]) -> list[InventoryLine]:
    """The inventory lines whose bills back an answer for this field."""
    if energy_type:
        return [ln for ln in lines if ln.energy_type.value == energy_type]
    if fld in ("scope1_tco2e", "fuel_consumption"):
        return [ln for ln in lines if ln.scope == 1]
    if fld in ("scope2_location_tco2e", "scope2_market_tco2e", "scope2_both", "electricity_mwh",
               "renewable_electricity_mwh", "nonrenewable_electricity_mwh"):
        return [ln for ln in lines if ln.scope == 2]
    return lines


def _bill_files(lines: list[InventoryLine], docs: dict[str, str]) -> list[str]:
    names = {docs.get(ln.source_doc_id) or f"document {ln.source_doc_id}" for ln in lines}
    return sorted(names)


# --- the pack -------------------------------------------------------------------------------------------

def _energy_types(lines: list[InventoryLine]) -> list[tuple[str, str]]:
    """(energy type, the unit its bills use) for each energy source in the inventory."""
    seen: dict[str, str] = {}
    for ln in lines:
        et = ln.energy_type.value
        if et == "refrigerant" or not (ln.scope == 1 or ln.basis == "location"):
            continue
        seen.setdefault(et, "kWh" if et == "electricity" else ln.unit.value)
    return sorted(seen.items())


def _source_name(et: str) -> str:
    return et.upper() if et in ("lpg", "cng", "lng") else et.replace("_", " ").capitalize()


def _cell(label: str | None, fld: str | None, ctx: FactContext, unit: str | None, energy_type: str | None = None,
          options: list[str] | None = None) -> PackCell:
    options = options or []
    if fld is None:
        return PackCell(label=label, unit=unit, portal=True, options=options,
                        reason="Answer this in the portal: it is a decision or a figure bills cannot give.")
    f = compute_fact(fld, ctx, unit=unit, energy_type=energy_type)
    if options:
        # a dropdown in the portal: the answer is one of its entries, or the supplier's to choose
        pick = None if f.flagged else choose_option(fld, f.value, options)
        if pick is None:
            why = ("your decision: a draft of the details is in the next question" if f.flagged and f.value is not None
                   else f.reason or "none of the portal's entries says what the data says")
            return PackCell(label=label, value=None, unit=unit, source=f.source, flagged=True, options=options,
                            reason=f"{why}. Choose from the portal's list.")
        return PackCell(label=label, value=pick, unit=unit, source=f.source, options=options)
    return PackCell(label=label, value=f.value, unit=unit, source=f.source, flagged=f.flagged or f.value is None, reason=f.reason)


def _question_cells(q: BankQuestion, ctx: FactContext) -> tuple[list[PackCell], list[tuple[str | None, str | None]]]:
    """The cells of one question, and the (field, energy type) pairs they came from, for the evidence."""
    cells: list[PackCell] = []
    origin: list[tuple[str | None, str | None]] = []
    if q.columns:
        for col in q.columns:
            cells.append(_cell(col.label, col.field, ctx, col.unit, options=q.options if col.field == q.field else None))
            origin.append((col.field, None))
    elif q.field == "energy_by_source" and q.format == "table":
        for et, unit in _energy_types(ctx.lines):
            c = _cell(_source_name(et), "energy_by_source", ctx, unit, et)
            cells.append(c)
            origin.append(("energy_by_source", et))
        if not cells:
            cells.append(PackCell(label=None, flagged=True, reason="no energy sources in this inventory", source="-"))
            origin.append(("energy_by_source", None))
    elif q.field == "energy_by_source" and q.format == "choice":
        types = [et.replace("_", " ") for et, _ in _energy_types(ctx.lines)]
        cells.append(PackCell(value=", ".join(types) or None, source="energy sources in this inventory", flagged=not types,
                              reason=None if types else "no energy sources in this inventory"))
        origin.append(("energy_by_source", None))
    else:
        cells.append(_cell(None, q.field, ctx, q.unit, options=q.options))
        origin.append((q.field, None))
    return cells, origin


def build_pack(bank: Bank | str, ctx: FactContext, *, docs: dict[str, str] | None = None) -> Pack:
    bank = get_bank(bank) if isinstance(bank, str) else bank
    docs = docs or {}
    items: list[PackItem] = []
    all_bills: set[str] = set()
    for q in bank.questions:
        cells, origin = _question_cells(q, ctx)
        computed = [c for c in cells if not c.portal]
        if not computed:
            status = "portal"
        elif any(c.flagged for c in computed):
            status = "needs_you"
        else:
            status = "answered"
        evidence: list[str] = []
        if status != "portal" or "bills" in q.evidence:
            if "bills" in q.evidence:
                files: set[str] = set()
                for fld, et in origin:
                    files.update(_bill_files(_relevant(fld, et, ctx.lines), docs))
                all_bills.update(files)
                evidence.append("Bills and invoices: " + (", ".join(sorted(files)) if files else "none uploaded for this period"))
            if "calc" in q.evidence:
                evidence.append(CALC_EVIDENCE)
            if "method" in q.evidence:
                evidence.append(METHOD_EVIDENCE)
        items.append(PackItem(
            ref=q.ref, ref_id=q.ref_id, text=q.text, wording=q.wording, format=q.format, status=status, cells=cells,
            evidence=evidence, note=q.note or q.options_note,
        ))
    p = ctx.profile
    return Pack(
        bank_id=bank.id, name=bank.name, publisher=bank.publisher, version=bank.version, wording=bank.wording,
        sources=bank.sources, notes=bank.notes, company=p.name or "", period=f"{p.fiscal_year_start:%d %b %Y} to {p.fiscal_year_end:%d %b %Y}",
        generated_at=datetime.now(UTC).strftime("%d %b %Y"), items=items, evidence_files=sorted(all_bills),
    )


# --- export -------------------------------------------------------------------------------------------------

def _shown(c: PackCell) -> str:
    if c.portal:
        return "Answer in the portal"
    if c.value is None:
        return "NOT PROVIDED"
    return f"{c.value:,}" if isinstance(c.value, int) else (f"{c.value:,.2f}".rstrip("0").rstrip(".") if isinstance(c.value, float) else str(c.value))


WORDING_NOTE = {
    "verbatim": "Question numbers and wording are the publisher's own.",
    "metric_name": "These are the metric names the publisher prints. EcoVadis does not publish question numbers or wording openly.",
    "topic": "These are our descriptions of the datapoints. The publisher does not publish the wording.",
}


def to_xlsx(pack: Pack, lines: list[InventoryLine], docs: dict[str, str] | None = None) -> bytes:
    docs = docs or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "Answers"
    bold = Font(bold=True)
    ws["A1"] = f"{pack.name}: answers to enter in the portal"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"{pack.company} · reporting period {pack.period} · prepared {pack.generated_at}"
    ws["A3"] = f"{pack.version}. {WORDING_NOTE[pack.wording]}"
    ws["A4"] = "Copy each answer into the portal and attach the evidence listed. Yellow rows need your confirmation first."
    head = ["Question", "Question text", "Answer", "Unit", "How it was computed", "Evidence to attach", "Status"]
    for i, h in enumerate(head, 1):
        c = ws.cell(6, i, h)
        c.font = bold
        c.fill = PatternFill("solid", fgColor="DDDDDD")
    r = 7
    for item in pack.items:
        for j, c in enumerate(item.cells):
            first = j == 0
            label = item.text if first else f"    ↳ {c.label or ''}"
            status = "Answer in the portal" if c.portal else ("Needs your answer" if c.flagged else "From your bills")
            row = [
                (item.ref + (f" ({item.ref_id})" if item.ref_id else "")) if first else None,
                label,
                _shown(c) if not (c.flagged and c.value is None and c.reason and not c.portal) else f"NOT PROVIDED: {c.reason}",
                c.unit or "",
                (c.source if not c.portal else (c.reason or "")) + (f" Portal options: {' / '.join(c.options)}." if c.options and (c.portal or c.value is None) else ""),
                "\n".join(item.evidence) if first else "",
                status,
            ]
            for k, v in enumerate(row, 1):
                cell = ws.cell(r, k, v)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.fill = PatternFill("solid", fgColor=GREY if c.portal else YELLOW if c.flagged else GREEN)
            if c.flagged and c.value is not None and c.reason:
                ws.cell(r, 5).value = f"{c.source}. Confirm: {c.reason}"
            r += 1
    for col, w in zip("ABCDEFG", (16, 60, 40, 12, 44, 48, 20), strict=True):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A7"

    calc = wb.create_sheet("Calculation lines")
    cols = ["Scope", "Basis", "Energy", "Quantity", "Unit", "Factor (kg CO2e per unit)", "Factor source", "kg CO2e", "Period", "Source document"]
    for i, h in enumerate(cols, 1):
        c = calc.cell(1, i, h)
        c.font = bold
        c.fill = PatternFill("solid", fgColor="DDDDDD")
    for n, ln in enumerate(lines, 2):
        vals = [ln.scope, ln.basis or "", ln.energy_type.value, ln.quantity, ln.unit.value, ln.factor_value, ln.factor_citation,
                round(ln.kg_co2e, 2), f"{ln.period_start} to {ln.period_end}", docs.get(ln.source_doc_id, ln.source_doc_id)]
        for k, v in enumerate(vals, 1):
            calc.cell(n, k, v)
    for i, w in enumerate((7, 10, 14, 12, 8, 16, 60, 12, 24, 36), 1):
        calc.column_dimensions[get_column_letter(i)].width = w

    ev = wb.create_sheet("Evidence to attach")
    ev["A1"] = "Documents to upload or have ready for the verifier"
    ev["A1"].font = bold
    for n, name in enumerate(pack.evidence_files or ["No bills uploaded for this period."], 2):
        ev.cell(n, 1, name)
    ev.cell(len(pack.evidence_files) + 3, 1, METHOD_EVIDENCE)
    ev.cell(len(pack.evidence_files) + 4, 1, CALC_EVIDENCE)
    ev.column_dimensions["A"].width = 90

    src = wb.create_sheet("Source")
    src["A1"] = "Where the questions come from"
    src["A1"].font = bold
    for n, s in enumerate(pack.sources, 2):
        src.cell(n, 1, s.title)
        src.cell(n, 2, s.url)
    for n, note in enumerate(pack.notes, len(pack.sources) + 3):
        src.cell(n, 1, note)
    src.column_dimensions["A"].width = 90
    src.column_dimensions["B"].width = 90

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def to_pdf(pack: Pack) -> bytes:
    html = env.get_template("pack.html").render(pack=pack, shown=_shown, wording_note=WORDING_NOTE[pack.wording], font_path=_font_path())
    return html_to_pdf(html)

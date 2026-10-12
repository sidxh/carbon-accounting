"""Score questionnaire filling against hand-worked answers.

    PYTHONPATH=. .venv/bin/python -m scripts.eval_questionnaires            # offline: harvest, answer and write
    PYTHONPATH=. .venv/bin/python -m scripts.eval_questionnaires --live     # also let Gemini map the wording
    PYTHONPATH=. .venv/bin/python -m scripts.eval_questionnaires --dir ~/my-sheets

Offline mode swaps the model for the mapping written in the corpus, so it checks everything that is code:
where each answer goes, unit conversion, dropdown choices, flags and the VSME named ranges. It is
deterministic and runs in CI. `--live` replaces the mapping with the model's, which is the one part that
can be wrong in a way code cannot predict. Responses are cached by the Gemini client.

The corpus is data/golden/questionnaires/corpus.json (see make_eval_corpus.py for what it holds). Every
expected value is worked out by hand from one fixed inventory. A case is one workbook and a list of
cells with what each must contain:

    value / startswith   the cell holds exactly this (numbers within 0.5%)
    flagged              the cell is yellow and holds no answer: it is the supplier's to fill
    unchanged            the buyer's own entry is still there

Two numbers matter most and are reported on their own. A wrong number is a cell that holds a different
figure than the arithmetic says. An unsafe fill is a cell that should have been left to the supplier but
was answered. Both must be zero.

Writes questionnaire_eval_report.json next to eval_report.json.
"""

from __future__ import annotations

import argparse
import io
import json
import warnings
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.agents.filler import fill_workbook
from app.facts import FactContext
from app.models import SupplierProfile
from app.vsme import fill_vsme

ROOT = Path(__file__).resolve().parents[2]
CORPUS = ROOT / "data" / "golden" / "questionnaires"

FLAG_RGB = "FFF2CC"


@dataclass
class CellResult:
    case: str
    where: str
    kind: str
    ok: bool
    got: Any = None
    expected: Any = None
    why: str = ""


@dataclass
class Report:
    cells: list[CellResult] = field(default_factory=list)
    cases: dict[str, dict] = field(default_factory=dict)

    @property
    def wrong_numbers(self) -> list[CellResult]:
        return [c for c in self.cells if c.kind == "number" and not c.ok and c.got is not None
                and not str(c.got).startswith("NOT PROVIDED")]

    @property
    def unsafe_fills(self) -> list[CellResult]:
        return [c for c in self.cells if c.kind == "flag" and not c.ok]

    @property
    def accuracy(self) -> float:
        return sum(c.ok for c in self.cells) / len(self.cells) if self.cells else 0.0


def _fixture(profile: dict[str, Any]) -> FactContext:
    from tests.qhelpers import ctx, vn_lines
    prof = SupplierProfile(id="eval", fiscal_year_start=date(2025, 1, 1), fiscal_year_end=date(2025, 12, 31), **profile)
    return ctx(vn_lines(), profile=prof)


def _close(got: Any, want: float) -> bool:
    try:
        return abs(float(got) - float(want)) <= max(0.005 * abs(float(want)), 0.005)
    except (TypeError, ValueError):
        return False


def _is_flagged(cell) -> bool:
    rgb = getattr(cell.fill.fgColor, "rgb", "") or ""
    return rgb.endswith(FLAG_RGB) and (cell.value is None or str(cell.value).startswith("NOT PROVIDED"))


def _check(case: str, spec: dict, cell, where: str) -> CellResult:
    kind = spec.get("kind", "number")
    got = cell.value
    if "unchanged" in spec:
        return CellResult(case, where, "unchanged", got == spec["unchanged"], got, spec["unchanged"], "the buyer's own entry was changed")
    if spec.get("flagged"):
        ok = _is_flagged(cell)
        return CellResult(case, where, "flag", ok, got, "left for the supplier", "" if ok else "answered a cell that is the supplier's to fill")
    if "startswith" in spec:
        ok = isinstance(got, str) and got.startswith(spec["startswith"])
        return CellResult(case, where, kind, ok, got if got is None else str(got)[:60], spec["startswith"] + "…")
    want = spec["value"]
    ok = _close(got, want) if isinstance(want, (int, float)) else got == want
    return CellResult(case, where, kind, ok, got, want)


def _excel_mapper(mapping: dict[str, list]):
    from app.agents.filler import Mapping, MappingResult

    def run(cells):
        out = []
        for c in cells:
            f, et = (mapping.get(c.text) or ["not_a_question", None])[:2]
            out.append(Mapping(question_id=c.id, field=f, energy_type=et, reasoning="corpus"))
        return MappingResult(mappings=out)
    return run


def run_excel(case: dict, base: Path, profile: dict, live: bool) -> list[CellResult]:
    path = (base / case["file"]).resolve()
    template = path.read_bytes()
    c = _fixture(profile)
    kw = {} if live else {"mapper": _excel_mapper(case["mapping"])}
    filled, _, _ = fill_workbook(
        template, profile=c.profile, snap=c.snap, lines=c.lines, assumptions=[], customer_share=0.2, customer_name="Toyota Boshoku", **kw,
    )
    wb = load_workbook(io.BytesIO(filled))
    return [_check(case["name"], s, wb[s["sheet"]][s["cell"]], f"{s['sheet']}!{s['cell']}") for s in case["expect"]]


def run_vsme(case: dict, profile: dict) -> list[CellResult]:
    res = fill_vsme(_fixture(profile))
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = load_workbook(io.BytesIO(res.xlsx))
    out = []
    for s in case["expect"]:
        sheet, ref = wb.defined_names[s["named_range"]].attr_text.rsplit("!", 1)
        cell = wb[sheet.strip("'")][ref.split(":")[0].replace("$", "")]
        out.append(_check(case["name"], s, cell, s["named_range"]))
    return out


def run_corpus(corpus_dir: Path, *, live: bool = False, only: str | None = None) -> Report:
    corpus = json.loads((corpus_dir / "corpus.json").read_text(encoding="utf-8"))
    report = Report()
    for case in corpus["cases"]:
        if only and only not in case["name"]:
            continue
        rows = run_vsme(case, corpus["profile"]) if case["kind"] == "vsme" else run_excel(case, corpus_dir, corpus["profile"], live)
        report.cells += rows
        report.cases[case["name"]] = {"lang": case["lang"], "shape": case["shape"], "cells": len(rows), "ok": sum(r.ok for r in rows)}
    return report


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--live", action="store_true", help="map wording with Gemini instead of the corpus mapping")
    ap.add_argument("--dir", type=Path, default=CORPUS, help="a directory with a corpus.json")
    ap.add_argument("--filter", help="substring on case name")
    args = ap.parse_args()

    report = run_corpus(args.dir, live=args.live, only=args.filter)
    for name, c in report.cases.items():
        mark = "ok " if c["ok"] == c["cells"] else "BAD"
        print(f"{mark} {name:32s} [{c['lang']}] {c['ok']}/{c['cells']}  {c['shape']}")
        for r in report.cells:
            if r.case == name and not r.ok:
                print(f"      {r.where}: got {r.got!r}, expected {r.expected!r} {r.why}")
    n = len(report.cells)
    print(f"\n{n} cells scored over {len(report.cases)} workbooks ({'live model mapping' if args.live else 'offline mapping'})")
    by_kind: dict[str, list[bool]] = {}
    for r in report.cells:
        by_kind.setdefault(r.kind, []).append(r.ok)
    for k, v in sorted(by_kind.items()):
        print(f"  {k:10s} {sum(v) / len(v):6.1%}  ({sum(v)}/{len(v)})")
    print(f"  overall    {report.accuracy:6.1%}")
    print(f"  wrong numbers {len(report.wrong_numbers)}    unsafe fills {len(report.unsafe_fills)}    (both must be 0)")

    (ROOT / "agent" / "questionnaire_eval_report.json").write_text(
        json.dumps({"live": args.live, "accuracy": report.accuracy, "cases": report.cases,
                    "cells": [r.__dict__ for r in report.cells]}, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
    ok = not report.wrong_numbers and not report.unsafe_fills and report.accuracy >= (0.9 if args.live else 1.0)
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())

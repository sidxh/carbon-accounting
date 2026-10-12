"""Build the offline CDP SME sheet from the official question bank.

    .venv/bin/python -m scripts.make_cdp_template

Writes data/templates/cdp_sme_questionnaire.xlsx. CDP itself is answered on the CDP Portal and CDP
publishes no Excel file for the SME questionnaire. Buyers still circulate their own sheet of the
same questions, so this is that sheet, laid out from `data/questionnaires/cdp_sme.yaml`: the same
numbers and the same wording CDP prints, nothing invented. A question that CDP asks as a table
becomes one row per column, under the numbered row, using CDP's own column headings.

Run it again whenever the bank changes; the test suite fails if the committed file is stale.
"""

from __future__ import annotations

from pathlib import Path
from typing import get_args

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from app.agents.filler import Field_
from app.questionnaires import get_bank

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "data" / "templates" / "cdp_sme_questionnaire.xlsx"

TITLE = "CDP SME climate questions — offline answer sheet"
LABEL = ("CDP itself is answered on the CDP Portal; this is the offline sheet buyers circulate. "
         "Question numbers and wording are CDP's own.")

# Which bank questions go on the sheet, in CDP's order, under which module heading.
SECTIONS = [
    ("Module 14: SME Introduction", ["14.4", "14.7"]),
    ("Module 20: SME Emissions", ["20.2", "20.4", "20.5", "20.7", "20.8", "20.9", "20.12", "20.15.2", "20.16", "20.17"]),
]

# Dropdowns come from the bank: the entries CDP prints in its own questionnaire PDF (20.8, 20.9, 20.17).
CHOICES = {q.ref: q.options for q in get_bank("cdp_sme").questions if q.options}

_FILLER_FIELDS = set(get_args(Field_))


def rows() -> list[tuple[str, str, str | None]]:
    """(ref, text, kind) for every line: kind 'section', 'q' (answered in its row), 'table' (answered below)
    or 'col' (one column of the table above)."""
    bank = get_bank("cdp_sme")
    out: list[tuple[str, str, str | None]] = []
    for heading, refs in SECTIONS:
        out.append(("", heading, "section"))
        for ref in refs:
            q = bank.question(ref)
            assert q is not None, ref
            cols = [c for c in q.columns if c.field in _FILLER_FIELDS and c.field != "not_a_question"]
            if len(cols) >= 2:
                out.append((ref, q.text, "table"))
                out.extend(("", c.label, "col") for c in cols)
            else:
                out.append((ref, q.text, "q"))
    return out


def build() -> tuple[Workbook, int]:
    """The workbook and the number of rows that take an answer."""
    bank = get_bank("cdp_sme")
    wb = Workbook()
    ws = wb.active
    ws.title = "Climate module"
    ws["A1"] = TITLE
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = LABEL
    ws["A2"].font = Font(italic=True, color="555555")
    ws["A3"] = f"Source: {bank.version}. {bank.sources[0].url}"
    ws["A3"].font = Font(size=9, color="777777")
    ws.append([])
    ws.append(["Ref", "Question", "Response", "Notes / evidence"])
    hdr = ws.max_row
    for c in "ABCD":
        ws[f"{c}{hdr}"].font = Font(bold=True)
        ws[f"{c}{hdr}"].fill = PatternFill("solid", fgColor="E8EEF5")
    lists = wb.create_sheet("Lists")
    lists.sheet_state = "hidden"        # the dropdown entries; a range, because one 20.9 entry contains a comma
    list_col: dict[str, str] = {}
    for i, (ref, entries) in enumerate(CHOICES.items(), start=1):
        for j, e in enumerate(entries, start=1):
            lists.cell(row=j, column=i, value=e)
        list_col[ref] = f"=Lists!${chr(64 + i)}$1:${chr(64 + i)}${len(entries)}"
    n = 0
    for ref, text, kind in rows():
        if kind == "section":
            ws.append([None, text, None, None])
            ws.cell(row=ws.max_row, column=2).font = Font(bold=True)
            ws.cell(row=ws.max_row, column=2).fill = PatternFill("solid", fgColor="F3F3F0")
            continue
        n += kind != "table"
        if kind == "table":
            ws.append([ref, text, "answer in the rows below", None])
            ws.cell(row=ws.max_row, column=3).font = Font(italic=True, color="888888")
        elif kind == "col":
            ws.append([None, text, None, None])
            ws.cell(row=ws.max_row, column=2).alignment = Alignment(indent=2, wrap_text=True, vertical="top")
        else:
            ws.append([ref, text, None, None])
            if ref in CHOICES:
                dv = DataValidation(type="list", formula1=list_col[ref], allow_blank=True)
                ws.add_data_validation(dv)
                dv.add(f"C{ws.max_row}")
    ws.column_dimensions["A"].width = 9
    ws.column_dimensions["B"].width = 84
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["D"].width = 50
    for row in ws.iter_rows(min_row=hdr + 1):
        for cell in row:
            if cell.alignment.indent == 0:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    return wb, n


def main() -> int:
    wb, n = build()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)} with {n} answer rows")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

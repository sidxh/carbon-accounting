"""Build the questionnaire evaluation corpus under data/golden/questionnaires/.

    PYTHONPATH=. .venv/bin/python scripts/make_eval_corpus.py

What it holds, and what is real:

* three bespoke buyer sheets (English, Japanese, Vietnamese). They are written by us, shaped like the
  sheets suppliers actually receive (a unit column, merged question cells, an answer box under a
  question, dropdowns from a range and inline). They are not copies of anyone's file.
* CDP SME: the offline sheet this repo generates from the official CDP bank.
* the official EFRAG VSME Digital Template (MIT licence), filled through its named ranges.

Not in the corpus: the JAPIA energy-use survey. JAPIA hands it to member companies as a zip and states no
licence that lets us republish it, so it stays out of the repository. `eval_questionnaires.py --dir` scores
any extra files you hold, using the same corpus.json format.

Every expected value is worked out by hand from one fixed inventory, so the evaluation checks the
product against arithmetic and not against itself:

    electricity  600,000 kWh location and 600,000 kWh market   -> 600 MWh
    LPG          10,000 kg at 2.94 kgCO2e/kg                   -> 29.4 t, and 131.39 MWh at 47.3 MJ/kg
    refrigerant  3.5 kg at GWP 2,088                           -> 7.308 t
    Scope 1 36.708 t    Scope 2 location 395.52 t    market 408 t    energy 731.39 MWh = 2,633 GJ
"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from app.questionnaires import get_bank

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "data" / "golden" / "questionnaires"

PROFILE = {"name": "Hoa Sen Garments", "country": "VN", "headcount": 420, "revenue": 24_000_000_000,
           "revenue_currency": "VND", "nace_code": "C.13.20"}

Mapping = dict[str, list]      # question text -> [field, energy_type or None]


def _save(wb: Workbook, name: str) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    wb.save(OUT / name)


def en_unit_column() -> dict:
    wb = Workbook()
    ws = wb.active
    ws.title = "Energy and emissions"
    ws["A1"] = "Supplier energy and emissions data request"
    ws["A2"] = "Please complete the Answer column for the last full financial year."
    ws.append([])
    ws.append(["Item", "Unit", "Answer", "Comment"])
    rows = [
        ("Name of company (completed by buyer)", None),
        ("Electricity purchased from the grid", "kWh"),
        ("LPG consumed", "t"),
        ("Diesel consumed", "L"),
        ("Total energy consumed", "MWh"),
        ("Scope 1 emissions", "t-CO2"),
        ("Scope 2 emissions, location-based", "kg-CO2"),
        ("Scope 2 emissions, market-based", "t-CO2"),
        ("Number of employees", None),
        ("Annual turnover", None),
        ("Are the figures verified by a third party?", None),
    ]
    for text, unit in rows:
        ws.append([text, unit, None, None])
    ws["C5"] = "Acme Co"                                    # the buyer already filled this one
    dv = DataValidation(type="list", formula1='"Yes,No"')
    ws.add_data_validation(dv)
    dv.add("C15")
    _save(wb, "bespoke_en_unit_column.xlsx")
    S = "Energy and emissions"
    mapping: Mapping = {
        "Electricity purchased from the grid": ["energy_by_source", "electricity"],
        "LPG consumed": ["energy_by_source", "lpg"],
        "Diesel consumed": ["energy_by_source", "diesel"],
        "Total energy consumed": ["total_energy_mwh", None],
        "Scope 1 emissions": ["scope1_tco2e", None],
        "Scope 2 emissions, location-based": ["scope2_location_tco2e", None],
        "Scope 2 emissions, market-based": ["scope2_market_tco2e", None],
        "Number of employees": ["employees", None],
        "Annual turnover": ["turnover", None],
        "Are the figures verified by a third party?": ["verification_status", None],
    }
    expect = [
        {"sheet": S, "cell": "C5", "unchanged": "Acme Co"},
        {"sheet": S, "cell": "C6", "value": 600_000, "kind": "number"},
        {"sheet": S, "cell": "C7", "value": 10, "kind": "number"},
        {"sheet": S, "cell": "C8", "flagged": True, "kind": "flag"},          # no diesel bill: the supplier decides
        {"sheet": S, "cell": "C9", "value": 731.4, "kind": "number"},
        {"sheet": S, "cell": "C10", "value": 36.71, "kind": "number"},
        {"sheet": S, "cell": "C11", "value": 395_520, "kind": "number"},      # the sheet asks for kg
        {"sheet": S, "cell": "C12", "value": 408, "kind": "number"},
        {"sheet": S, "cell": "C13", "value": 420, "kind": "number"},
        {"sheet": S, "cell": "C14", "value": 24_000_000_000, "kind": "number"},
        {"sheet": S, "cell": "C15", "value": "No", "kind": "choice"},
    ]
    return {"name": "bespoke_en_unit_column", "kind": "excel", "file": "bespoke_en_unit_column.xlsx", "lang": "en",
            "shape": "header row with unit and answer columns, dropdown, pre-answered cell", "mapping": mapping, "expect": expect}


def ja_energy_table() -> dict:
    wb = Workbook()
    ws = wb.active
    ws.title = "調査票"
    ws["A1"] = "エネルギー使用実績調査票(当社作成のサンプル)"
    ws.append([])
    ws.append(["項目", "単位", "回答", "備考"])
    rows = [
        ("購入電力量", "kWh"), ("LPG使用量", "t"), ("軽油使用量", "kL"), ("エネルギー使用量合計", "GJ"),
        ("CO2排出量(スコープ1)", "t-CO2"), ("CO2排出量(スコープ2・ロケーション基準)", "t-CO2"),
        ("所在国", None), ("従業員数", "人"),
    ]
    for text, unit in rows:
        ws.append([text, unit, None, None])
    lists = wb.create_sheet("リスト")
    for i, v in enumerate(("日本", "ベトナム", "タイ"), 1):
        lists.cell(i, 1, v)
    lists.sheet_state = "hidden"
    dv = DataValidation(type="list", formula1="リスト!$A$1:$A$3")
    ws.add_data_validation(dv)
    dv.add("C10")
    _save(wb, "bespoke_ja_energy_table.xlsx")
    S = "調査票"
    mapping: Mapping = {
        "購入電力量": ["energy_by_source", "electricity"], "LPG使用量": ["energy_by_source", "lpg"],
        "軽油使用量": ["energy_by_source", "diesel"], "エネルギー使用量合計": ["total_energy_mwh", None],
        "CO2排出量(スコープ1)": ["scope1_tco2e", None], "CO2排出量(スコープ2・ロケーション基準)": ["scope2_location_tco2e", None],
        "所在国": ["country", None], "従業員数": ["employees", None],
    }
    expect = [
        {"sheet": S, "cell": "C4", "value": 600_000, "kind": "number"},
        {"sheet": S, "cell": "C5", "value": 10, "kind": "number"},
        {"sheet": S, "cell": "C6", "flagged": True, "kind": "flag"},
        {"sheet": S, "cell": "C7", "value": 2633, "kind": "number"},          # 731.39 MWh x 3.6
        {"sheet": S, "cell": "C8", "value": 36.71, "kind": "number"},
        {"sheet": S, "cell": "C9", "value": 395.52, "kind": "number"},
        {"sheet": S, "cell": "C10", "value": "ベトナム", "kind": "choice"},
        {"sheet": S, "cell": "C11", "value": 420, "kind": "number"},
    ]
    return {"name": "bespoke_ja_energy_table", "kind": "excel", "file": "bespoke_ja_energy_table.xlsx", "lang": "ja",
            "shape": "Japanese headers, unit column in kL and GJ, dropdown from a hidden list sheet", "mapping": mapping, "expect": expect}


def vi_merged() -> dict:
    wb = Workbook()
    ws = wb.active
    ws.title = "Khảo sát"
    ws["A1"] = "Phiếu khảo sát phát thải của nhà cung cấp"
    ws.append([])
    ws["A3"], ws["B3"] = "1", "Tổng lượng phát thải Phạm vi 1 trong năm báo cáo (tCO2e)?"
    ws.merge_cells("B3:D3")
    ws["A4"], ws["B4"] = "2", "Tổng lượng phát thải Phạm vi 2 theo phương pháp thị trường (tCO2e)?"
    ws.merge_cells("B4:D4")
    ws["A5"], ws["B5"] = "3", "Mô tả phương pháp doanh nghiệp dùng để tính phát thải."
    ws.merge_cells("B5:E5")
    ws.merge_cells("B6:E6")                                  # the answer box under question 3
    ws["A7"], ws["B7"] = "4", "Số liệu phát thải có được bên thứ ba kiểm chứng không?"
    ws.merge_cells("B7:D7")
    dv = DataValidation(type="list", formula1='"Có,Không"')
    ws.add_data_validation(dv)
    dv.add("E7")
    _save(wb, "bespoke_vi_merged.xlsx")
    S = "Khảo sát"
    mapping: Mapping = {
        "Tổng lượng phát thải Phạm vi 1 trong năm báo cáo (tCO2e)?": ["scope1_tco2e", None],
        "Tổng lượng phát thải Phạm vi 2 theo phương pháp thị trường (tCO2e)?": ["scope2_market_tco2e", None],
        "Mô tả phương pháp doanh nghiệp dùng để tính phát thải.": ["methodology", None],
        "Số liệu phát thải có được bên thứ ba kiểm chứng không?": ["verification_status", None],
    }
    expect = [
        {"sheet": S, "cell": "E3", "value": 36.71, "kind": "number"},
        {"sheet": S, "cell": "E4", "value": 408, "kind": "number"},
        {"sheet": S, "cell": "B6", "startswith": "GHG Protocol Corporate Standard", "kind": "text"},
        {"sheet": S, "cell": "E7", "value": "Không", "kind": "choice"},
    ]
    return {"name": "bespoke_vi_merged", "kind": "excel", "file": "bespoke_vi_merged.xlsx", "lang": "vi",
            "shape": "merged question cells, answer right of the merge, answer box below a question, inline dropdown", "mapping": mapping, "expect": expect}


def cdp_offline() -> dict:
    """Our offline copy of the CDP SME climate sheet, with every expectation worked out above."""
    path = ROOT / "data" / "templates" / "cdp_sme_questionnaire.xlsx"
    ws = load_workbook(path)["Climate module"]
    row_of = {}
    for r in ws.iter_rows(min_row=6):
        if r[1].value:
            row_of[str(r[1].value)] = r[1].row
    bank = get_bank("cdp_sme")
    mapping: Mapping = {}
    for q in bank.questions:
        mapping.setdefault(q.text, [q.field or "not_a_question", None])
        for c in q.columns:
            mapping.setdefault(c.label, [("methodology" if c.label == "Methodological details" else c.field) or "not_a_question", None])

    def at(text: str) -> str:
        return f"C{row_of[text]}"

    S = "Climate module"
    expect = [
        {"sheet": S, "cell": at("State the end date of the year for which you are reporting data."), "value": "2025-12-31", "kind": "text"},
        {"sheet": S, "cell": at("Gross global Scope 1 emissions (metric tons CO2e)"), "value": 36.71, "kind": "number"},
        {"sheet": S, "cell": at("Gross global Scope 2, location-based emissions (metric tons CO2e)"), "value": 395.52, "kind": "number"},
        {"sheet": S, "cell": at("Gross global Scope 2, market-based emissions (metric tons CO2e)"), "value": 408, "kind": "number"},
        {"sheet": S, "cell": at("Indicate the verification/assurance status that applies to your reported emissions."),
         "value": "No third-party verification or assurance", "kind": "choice"},
        # a first reporting year in this account: the supplier states whether it really is the first
        {"sheet": S, "cell": at("How do your gross global emissions (Scope 1 and 2 combined) for the reporting year compare to those of the previous reporting year?"),
         "flagged": True, "kind": "flag"},
        {"sheet": S, "cell": at("MWh from renewable sources"), "flagged": True, "kind": "flag"},       # a green tariff is never assumed
        {"sheet": S, "cell": at("Total (renewable + non-renewable) MWh"), "value": 600, "kind": "number"},
    ]
    return {"name": "cdp_sme_offline_sheet", "kind": "excel", "file": "../../templates/cdp_sme_questionnaire.xlsx", "lang": "en",
            "shape": "CDP numbering, table questions as one row per column, dropdowns from a hidden list", "mapping": mapping, "expect": expect}


def efrag_vsme() -> dict:
    """EFRAG's official template, filled through named ranges, so nothing is mapped by a model."""
    return {
        "name": "efrag_vsme_digital_template", "kind": "vsme", "lang": "en",
        "shape": "official EFRAG xlsx, named ranges, dropdown validations kept",
        "expect": [
            {"named_range": "template_reporting_entity_name", "value": "Hoa Sen Garments", "kind": "text"},
            {"named_range": "template_currency", "value": "VND", "kind": "choice"},
            {"named_range": "NaceSectorClassificationCodes", "value": "NACE C - 13.20 Weaving of textiles", "kind": "choice"},
            {"named_range": "Turnover", "value": 24_000_000_000, "kind": "number"},
            {"named_range": "NumberOfEmployees", "value": 420, "kind": "number"},
            {"named_range": "CountryOfPrimaryOperationsAndLocationOfSignificantAssets", "value": "Viet Nam", "kind": "choice"},
            {"named_range": "TotalEnergyConsumption", "value": 731.4, "kind": "number"},
            {"named_range": "GrossScope1GreenhouseGasEmissions", "value": 36.71, "kind": "number"},
            {"named_range": "GrossLocationBasedScope2GreenhouseGasEmissions", "value": 395.52, "kind": "number"},
            {"named_range": "GrossMarketBasedScope2GreenhouseGasEmissions", "value": 408, "kind": "number"},
        ],
    }


def main() -> None:
    cases = [en_unit_column(), ja_energy_table(), vi_merged(), cdp_offline(), efrag_vsme()]
    (OUT / "corpus.json").write_text(json.dumps({"profile": PROFILE, "cases": cases}, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"wrote {len(cases)} cases to {OUT}")


if __name__ == "__main__":
    main()

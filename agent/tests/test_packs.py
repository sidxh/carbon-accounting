"""Answer packs: every question in the official bank, answered from the same facts as the Excel filler."""

import io

import pytest
from openpyxl import load_workbook

from app.facts import PreviousYear, compute_fact
from app.packs import PACK_BANKS, build_pack, to_pdf, to_xlsx
from app.questionnaires import get_bank

from .qhelpers import ctx, profile, vn_lines

DOCS = {"d": "EVN-2025-01.pdf"}
FULL = {"headcount": 420, "revenue": 24_000_000_000, "revenue_currency": "VND", "nace_code": "C.13.20"}


def _pack(bank, **kw):
    return build_pack(bank, ctx(profile=profile(**FULL), customer_share=0.2, customer_name="Toyota Boshoku", **kw), docs=DOCS)


@pytest.mark.parametrize("bank", PACK_BANKS)
def test_a_pack_has_every_question_of_its_bank_in_order(bank):
    p = _pack(bank)
    assert [i.ref for i in p.items] == get_bank(bank).refs
    assert p.version == get_bank(bank).version and p.sources
    assert sum(p.counts.values()) == len(p.items)


def test_cdp_pack_uses_cdps_numbering_and_the_calculators_numbers():
    p = {i.ref: i for i in _pack("cdp_sme").items}
    assert p["20.4"].cells[0].value == 36.71 and p["20.4"].cells[0].unit == "tCO2e"
    five = p["20.5"]
    assert [c.value for c in five.cells[:2]] == [395.52, 408.0]           # location and market, one question with columns
    assert five.cells[2].label == "Methodological details"
    assert p["14.4"].cells[0].value == "2025-12-31"
    assert p["20.12"].cells[0].value == "Toyota Boshoku"
    assert p["20.12"].cells[5].value == pytest.approx(round((36.708 + 395.52) * 0.2, 2), abs=0.02)


def test_the_pack_and_the_filler_give_the_same_number():
    c = ctx(profile=profile(**FULL))
    assert _pack("cdp_sme").items[5].cells[0].value == compute_fact("scope1_tco2e", c).value
    assert _pack("ecovadis_carbon").items[0].cells[0].value == compute_fact("scope1_tco2e", c).value


def test_choice_questions_answer_with_the_portals_own_entry():
    p = {i.ref: i for i in _pack("cdp_sme").items}
    assert p["20.8"].cells[0].value == "No third-party verification or assurance"
    assert p["20.9"].status == "needs_you" and p["20.9"].cells[0].value is None            # first year: only the supplier knows
    prev = PreviousYear("2024", scope1_kg=30_000, scope2_location_kg=390_000)
    p2 = {i.ref: i for i in _pack("cdp_sme", previous=prev).items}
    assert p2["20.9"].cells[0].value == "Increased" and p2["20.9"].status == "answered"
    assert p["20.17"].cells[0].value is None and "Yes" in p["20.17"].cells[0].options       # a decision, never guessed


def test_what_only_the_supplier_can_say_is_marked_for_the_portal():
    p = {i.ref: i for i in _pack("cdp_sme").items}
    assert p["20.1"].status == "portal" and p["20.6"].status == "portal"
    assert all(c.portal and c.value is None for c in p["20.1"].cells)
    assert p["20.15.2"].status == "needs_you"                                # renewable MWh is never assumed
    renewable = p["20.15.2"].cells[0]
    assert renewable.value is None and renewable.flagged and "green tariff" in renewable.reason


def test_higg_pack_lists_energy_sources_with_the_unit_their_bills_use():
    p = {i.ref: i for i in _pack("higg_fem_energy").items}
    assert p["Q1"].cells[0].value == "electricity, lpg"
    rows = {c.label: (c.value, c.unit) for c in p["Q1 quantity"].cells}
    assert rows == {"Electricity": (600_000, "kWh"), "LPG": (10_000, "kg")}
    assert p["Q2"].status == "portal"                                         # owned vehicles: only the supplier knows


def test_evidence_names_the_uploaded_bills_and_the_derived_documents():
    items = {i.ref: i for i in _pack("cdp_sme").items}
    assert items["20.4"].evidence[0] == "Bills and invoices: EVN-2025-01.pdf"
    assert any("Calculation lines" in e for e in items["20.4"].evidence)
    assert any("Methodology note" in e for e in items["20.2"].evidence)
    assert items["20.1"].evidence == []
    assert _pack("cdp_sme").evidence_files == ["EVN-2025-01.pdf"]
    assert build_pack("cdp_sme", ctx(), docs={}).items[5].evidence[0] == "Bills and invoices: document d"      # a deleted file is still named
    assert build_pack("cdp_sme", ctx([]), docs={}).items[5].evidence[0].endswith("none uploaded for this period")


def test_flagged_answers_are_never_shown_as_final():
    p = build_pack("cdp_sme", ctx())                                         # no headcount, no revenue
    assert p.counts["needs_you"] >= 4
    for i in p.items:
        for c in i.cells:
            if c.flagged:
                assert c.reason


def test_xlsx_export_has_the_answers_the_calculation_lines_and_the_evidence():
    p = _pack("cdp_sme")
    wb = load_workbook(io.BytesIO(to_xlsx(p, vn_lines(), DOCS)))
    assert wb.sheetnames == ["Answers", "Calculation lines", "Evidence to attach", "Source"]
    ws = wb["Answers"]
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=7)]
    row = next(r for r in rows if r[0] == "20.4")
    assert row[1].startswith("What were your organization") and row[2] == "36.71" and row[3] == "tCO2e"
    assert "EVN-2025-01.pdf" in row[5]
    portal = next(r for r in rows if r[0] == "20.1")
    assert portal[2] == "Answer in the portal"
    assert wb["Calculation lines"].max_row == 1 + len(vn_lines())
    assert wb["Evidence to attach"]["A2"].value == "EVN-2025-01.pdf"
    fills = {ws.cell(i, 1).value: ws.cell(i, 1).fill.fgColor.rgb[-6:] for i in range(7, ws.max_row + 1) if ws.cell(i, 1).value}
    assert fills["20.4"] == "E3F1E4" and fills["20.1"] == "ECECEC" and fills["20.15.2"] == "FFF2CC"


def test_pdf_export_renders():
    pdf = to_pdf(_pack("ecovadis_carbon"))
    assert pdf[:5] == b"%PDF-" and len(pdf) > 2_000

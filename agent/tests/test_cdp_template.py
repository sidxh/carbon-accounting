"""The offline CDP sheet is built from the official bank, and the filler answers it correctly end to end."""

import io
from pathlib import Path

from openpyxl import load_workbook

from app.agents.filler import Mapping, MappingResult, answer_all, write_answers
from app.agents.harvest import harvest
from app.questionnaires import get_bank
from scripts.make_cdp_template import build

from .qhelpers import ctx, vn_lines

TEMPLATE = Path(__file__).resolve().parents[2] / "data" / "templates" / "cdp_sme_questionnaire.xlsx"


def _values(wb):
    return [[c.value for c in row] for row in wb["Climate module"].iter_rows()]


def test_committed_sheet_is_what_the_bank_generates():
    fresh, _ = build()
    assert _values(fresh) == _values(load_workbook(TEMPLATE)), "run `python -m scripts.make_cdp_template`"


def test_sheet_is_labelled_as_the_offline_copy_and_uses_official_numbers():
    ws = load_workbook(TEMPLATE)["Climate module"]
    assert "CDP itself is answered on the CDP Portal" in ws["A2"].value
    refs = {str(r[0].value) for r in ws.iter_rows(min_row=6) if r[0].value}
    assert {"14.4", "14.7", "20.2", "20.4", "20.5", "20.7", "20.8", "20.9", "20.12", "20.15.2", "20.16", "20.17"} <= refs
    assert not refs & {"20.5a", "20.5b", "20.12a", "20.12b", "20.15.3", "0.1", "DQ.1"}
    # every wording on the sheet is one the bank holds
    bank = get_bank("cdp_sme")
    bank_text = {q.text for q in bank.questions} | {c.label for q in bank.questions for c in q.columns}
    for row in ws.iter_rows(min_row=6):
        if row[0].value:
            assert row[1].value in bank_text, row[1].value


def _bank_mapper(cells):
    """What a perfect model would answer for CDP's own wording, read off the bank."""
    bank = get_bank("cdp_sme")
    by_text = {}
    for q in bank.questions:
        by_text[q.text] = q.field
        for c in q.columns:
            by_text.setdefault(c.label, c.field)
    out = []
    for c in cells:
        if c.text.startswith("Module "):
            out.append(Mapping(question_id=c.id, field="not_a_question", reasoning="heading"))
            continue
        f = by_text[c.text]
        if c.text == "Methodological details":
            f = "methodology"
        out.append(Mapping(question_id=c.id, field=f, reasoning="bank"))
    return MappingResult(mappings=out)


def _fill(**kw):
    wb = load_workbook(TEMPLATE)
    cells = harvest(wb)
    answers = answer_all(cells, _bank_mapper(cells), profile=ctx().profile, snap=ctx().snap, lines=vn_lines(), assumptions=[], **kw)
    out = load_workbook(io.BytesIO(write_answers(TEMPLATE.read_bytes(), answers)))["Climate module"]
    return {a.ref or a.question: a for a in answers}, out, answers


def test_every_cdp_row_lands_in_its_own_response_cell():
    _, ws, answers = _fill()
    by_q = {a.question: a for a in answers}
    s1 = by_q["Gross global Scope 1 emissions (metric tons CO2e)"]
    assert ws.cell(s1.answer_row, 3).value == round((10_000 * 2.94 + 3.5 * 2088) / 1000, 2)       # LPG plus the refrigerant top-up
    loc = by_q["Gross global Scope 2, location-based emissions (metric tons CO2e)"]
    mkt = by_q["Gross global Scope 2, market-based emissions (metric tons CO2e)"]
    assert ws.cell(loc.answer_row, 3).value == round(600_000 * 0.6592 / 1000, 2)
    assert ws.cell(mkt.answer_row, 3).value == round(600_000 * 0.68 / 1000, 2)
    assert mkt.answer_row == loc.answer_row + 1                       # one row per column of 20.5, in order
    total = by_q["Total (renewable + non-renewable) MWh"]
    assert ws.cell(total.answer_row, 3).value == 600.0               # 20.15.2


def test_cdp_dropdowns_get_an_entry_or_stay_open():
    _, ws, answers = _fill()
    by_ref = {a.ref: a for a in answers if a.ref}
    # verification is "no" by construction, so its entry is picked and the full statement goes to the notes
    v = by_ref["20.8"]
    assert ws.cell(v.answer_row, 3).value == "No third-party verification or assurance"
    assert ws.cell(v.answer_row, 4).value.startswith("Not verified.")
    # last year is unknown in a first year, so 20.9 stays empty and yellow rather than guessing an entry
    y = by_ref["20.9"]
    assert ws.cell(y.answer_row, 3).value is None and y.flagged
    assert ws.cell(y.answer_row, 3).fill.fgColor.rgb.endswith("FFF2CC")
    # a plan to cut emissions is the supplier's to state: no entry is chosen for them
    r = by_ref["20.17"]
    assert ws.cell(r.answer_row, 3).value is None and r.flagged


def test_cdp_first_year_answered_when_a_previous_year_exists():
    from app.facts import PreviousYear

    prev = PreviousYear("2024-01-01 to 2024-12-31", scope1_kg=40_000, scope2_location_kg=420_000)
    _, ws, answers = _fill(previous=prev)
    y = next(a for a in answers if a.ref == "20.9")
    assert ws.cell(y.answer_row, 3).value == "Decreased"
    assert "Decreased." in y.note and "2024-01-01 to 2024-12-31" in y.note

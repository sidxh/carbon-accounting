"""The questionnaire evaluation itself: the committed corpus is current, it passes, and it can fail."""

import json
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from scripts import make_eval_corpus
from scripts.eval_questionnaires import CORPUS, run_corpus


def test_corpus_passes_offline_with_no_wrong_number_and_no_unsafe_fill():
    r = run_corpus(CORPUS)
    assert len(r.cells) >= 40 and len(r.cases) == 5
    assert [c for c in r.cells if not c.ok] == []
    assert r.wrong_numbers == [] and r.unsafe_fills == []


def test_committed_corpus_is_what_the_builder_writes(tmp_path, monkeypatch):
    monkeypatch.setattr(make_eval_corpus, "OUT", tmp_path)
    make_eval_corpus.main()
    assert json.loads((tmp_path / "corpus.json").read_text("utf-8")) == json.loads((CORPUS / "corpus.json").read_text("utf-8"))
    for f in tmp_path.glob("*.xlsx"):
        a, b = load_workbook(f), load_workbook(CORPUS / f.name)
        assert [[c.value for c in r] for ws in a for r in ws.iter_rows()] == [[c.value for c in r] for ws in b for r in ws.iter_rows()], f.name


@pytest.fixture
def copy(tmp_path: Path):
    d = tmp_path / "golden" / "questionnaires"
    shutil.copytree(CORPUS, d)
    (tmp_path / "templates").mkdir()
    shutil.copy(CORPUS.parents[1] / "templates" / "cdp_sme_questionnaire.xlsx", tmp_path / "templates")
    shutil.copytree(CORPUS.parents[1] / "templates" / "vsme", tmp_path / "templates" / "vsme")
    return d


def _edit(d: Path, case: str, cell: str, **new):
    data = json.loads((d / "corpus.json").read_text("utf-8"))
    spec = next(s for c in data["cases"] if c["name"] == case for s in c["expect"] if s.get("cell") == cell or s.get("named_range") == cell)
    for k in [k for k in spec if k in ("value", "flagged", "startswith", "unchanged")]:
        del spec[k]
    spec.update(new)
    (d / "corpus.json").write_text(json.dumps(data, ensure_ascii=False), "utf-8")


def test_a_wrong_expected_number_is_reported_as_a_wrong_number(copy):
    _edit(copy, "bespoke_en_unit_column", "C10", value=99.0)
    r = run_corpus(copy)
    assert [c.where for c in r.wrong_numbers] == ["Energy and emissions!C10"]


def test_an_answered_cell_that_should_have_been_left_open_is_an_unsafe_fill(copy):
    _edit(copy, "bespoke_en_unit_column", "C6", flagged=True, kind="flag")      # electricity is on the bills, so answering it is "unsafe" here
    r = run_corpus(copy)
    assert [c.where for c in r.unsafe_fills] == ["Energy and emissions!C6"]


def test_a_missing_mapping_leaves_the_cell_empty_which_scores_as_a_miss(copy):
    data = json.loads((copy / "corpus.json").read_text("utf-8"))
    del next(c for c in data["cases"] if c["name"] == "bespoke_ja_energy_table")["mapping"]["購入電力量"]
    (copy / "corpus.json").write_text(json.dumps(data, ensure_ascii=False), "utf-8")
    r = run_corpus(copy, only="ja_energy")
    assert [c.where for c in r.cells if not c.ok] == ["調査票!C4"]

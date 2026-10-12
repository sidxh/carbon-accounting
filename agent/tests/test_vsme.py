"""The official EFRAG VSME Digital Template, filled through its named ranges with no model involved."""

import io
import re
import warnings
import zipfile

import pytest
from openpyxl import load_workbook

from app.questionnaires import get_bank, templates_dir
from app.vsme import VsmeError, _patch, fill_vsme, nace_entry

from .qhelpers import ctx, profile

TEMPLATE = templates_dir() / get_bank("vsme_b3").template_file
FULL = {"headcount": 420, "revenue": 24_000_000_000, "revenue_currency": "VND", "nace_code": "C.13.20"}


def _open(xlsx: bytes):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return load_workbook(io.BytesIO(xlsx))


@pytest.fixture(scope="module")
def filled():
    return fill_vsme(ctx(profile=profile(**FULL)))


def test_values_land_in_the_cells_the_named_ranges_point_to(filled):
    wb = _open(filled.xlsx)
    gi, env = wb["General Information"], wb["Environmental Disclosures"]

    def at(name: str, dx=0, dy=0):
        sheet, ref = wb.defined_names[name].attr_text.rsplit("!", 1)
        c = wb[sheet.strip("'")][ref.split(":")[0].replace("$", "")]
        return wb[sheet.strip("'")].cell(c.row + dy, c.column + dx).value

    assert at("template_reporting_entity_name") == "Hoa Sen Garments"
    assert at("template_currency") == "VND"
    assert at("NaceSectorClassificationCodes") == "NACE C - 13.20 Weaving of textiles"
    assert at("Turnover") == 24_000_000_000
    assert at("NumberOfEmployees") == 420 and at("NumberOfEmployees", dy=2) == "Headcount"
    assert at("CountryOfPrimaryOperationsAndLocationOfSignificantAssets") == "Viet Nam"
    assert at("TotalEnergyConsumption") == 731.4
    assert at("GrossScope1GreenhouseGasEmissions") == 36.71
    assert at("GrossLocationBasedScope2GreenhouseGasEmissions") == 395.52
    assert at("GrossMarketBasedScope2GreenhouseGasEmissions") == 408.0
    # the date range is a formula over three input cells: year, month, day
    assert [gi[f"D{r}"].value for r in (6, 7, 8)] == [2025, 1, 1] and [gi[f"D{r}"].value for r in (10, 11, 12)] == [2025, 12, 31]
    assert gi["D9"].value.startswith("=IF(") and gi["D13"].value.startswith("=IF(")
    assert env["D25"].value.startswith("=IF(")            # the total row is still a formula


def test_every_filled_input_is_one_the_templates_own_validation_reads(filled):
    """The template has a status cell per input (F3, F5 ...). Each status formula reads the input it checks, so
    an input we fill that its status formula does not mention would be in the wrong place."""
    wb = _open(filled.xlsx)
    checked = 0
    for f in filled.filled:
        sheet, coord = f.cell.rsplit("!", 1)
        sheet = sheet.strip("'")
        _, row = re.fullmatch(r"([A-Z]+)(\d+)", coord).groups()
        if sheet == "General Information" and row in {"3", "5", "182", "283", "284", "287"}:
            assert re.search(rf"\b{coord}\b", str(wb[sheet][f"F{row}"].value)), coord
            checked += 1
        if coord == "G5":
            assert re.search(r"\bG5\b", str(wb[sheet]["L5"].value))
            checked += 1
        if sheet == "Environmental Disclosures" and coord in {"D22", "D23", "D24"}:
            assert re.search(rf"\b{coord}\b", str(wb[sheet][f"L{row}"].value)), coord
            checked += 1
    assert checked >= 10
    gi = wb["General Information"]
    assert all(c in gi["F6"].value for c in ("D6", "D7", "D8")) and all(c in gi["F10"].value for c in ("D10", "D11", "D12"))


def test_the_rest_of_the_package_is_untouched(filled):
    with zipfile.ZipFile(TEMPLATE) as a, zipfile.ZipFile(io.BytesIO(filled.xlsx)) as b:
        assert [i.filename for i in a.infolist()] == [i.filename for i in b.infolist()]
        changed = {n for n in a.namelist() if a.read(n) != b.read(n)}
        sheet_gi = "xl/worksheets/sheet4.xml"
        assert changed <= {"xl/workbook.xml", sheet_gi, "xl/worksheets/sheet5.xml"}, changed
        # the dropdown extensions openpyxl would have dropped are all still there
        assert b.read(sheet_gi).count(b"<x14:dataValidation ") == a.read(sheet_gi).count(b"<x14:dataValidation ") > 0
        assert b.read(sheet_gi).count(b"<f>") == a.read(sheet_gi).count(b"<f>")      # no formula was overwritten
        assert b"fullCalcOnLoad" in b.read("xl/workbook.xml")


def test_what_it_could_not_fill_is_reported_not_guessed(filled):
    refs = {o.ref for o in filled.open_items}
    assert {"B1 4", "B1 basis", "B3 29 breakdown", "B1 24(e)(v)"} <= refs


def test_missing_profile_data_is_left_empty_and_listed():
    r = fill_vsme(ctx(profile=profile()))
    wb = _open(r.xlsx)
    assert wb["General Information"]["E284"].value is None and wb["General Information"]["E283"].value is None
    assert wb["General Information"]["E182"].value is None
    reasons = " ".join(o.reason for o in r.open_items)
    assert "number of employees" in reasons and "NACE" in reasons
    assert wb["Environmental Disclosures"]["D22"].value == 36.71          # the emissions still go in


def test_a_value_that_is_not_in_the_templates_list_is_not_written():
    r = fill_vsme(ctx(profile=profile(country="ZZ", nace_code="Q.99.99", revenue_currency="XXX", revenue=1000)))
    wb = _open(r.xlsx)
    gi = wb["General Information"]
    assert gi["E287"].value is None and gi["E182"].value is None and gi["D5"].value is None
    assert sum("not in the template" in o.reason for o in r.open_items) >= 3


def test_total_energy_with_an_unconvertible_fuel_is_left_for_the_supplier():
    from app.models import EnergyType, Unit

    from .qhelpers import inv, vn_lines
    r = fill_vsme(ctx(vn_lines() + [inv(EnergyType.COAL, 100, Unit.KG, 1, scope=1)], profile=profile(**FULL)))
    assert _open(r.xlsx)["Environmental Disclosures"]["G5"].value is None
    assert any(o.ref == "B3 29 total" and "coal" in o.reason for o in r.open_items)


def test_a_cell_holding_a_formula_is_never_overwritten():
    with pytest.raises(VsmeError):
        _patch('<row><c r="D9" s="1" t="str"><f>A1</f><v>-</v></c></row>', "D9", 5)
    assert _patch('<c r="D3" s="646"/>', "D3", "A & B") == '<c r="D3" s="646" t="inlineStr"><is><t xml:space="preserve">A &amp; B</t></is></c>'
    assert _patch('<c r="D30" s="1"/><c r="D3" s="2"/>', "D3", 7) == '<c r="D30" s="1"/><c r="D3" s="2"><v>7</v></c>'


def test_nace_codes_match_the_templates_list_text():
    entries = ["NACE C - MANUFACTURING", "NACE C - 13 Manufacture of textiles", "NACE C - 13.2 Weaving", "NACE C - 13.20 Weaving of textiles",
               "NACE G - 47 Retail trade"]
    assert nace_entry("C.13.20", entries) == "NACE C - 13.20 Weaving of textiles"
    assert nace_entry("13.20", entries) == "NACE C - 13.20 Weaving of textiles"
    assert nace_entry("NACE C - 13", entries) == "NACE C - 13 Manufacture of textiles"
    assert nace_entry("c 13.2", entries) == "NACE C - 13.2 Weaving"
    assert nace_entry("C", entries) == "NACE C - MANUFACTURING"
    assert nace_entry("G.13.20", entries) is None and nace_entry("", entries) is None and nace_entry("hello", entries) is None


def test_the_bank_names_every_range_the_template_has():
    names = set(re.findall(r'<definedName name="([^"]+)"', zipfile.ZipFile(TEMPLATE).read("xl/workbook.xml").decode()))
    for q in get_bank("vsme_b3").questions:
        assert q.named_range in names, q.ref

"""VSME and answer-pack routes against the in-memory repo, offline."""

import io
import warnings
import zipfile
from datetime import date

from openpyxl import load_workbook

from app.models import ActivityLine, EnergyType, Unit

from .test_api import _auth, _setup

PROFILE = {
    "fiscal_year_start": "2025-01-01", "fiscal_year_end": "2025-12-31", "country": "VN", "name": "Hoa Sen Garments",
    "headcount": 420, "revenue": 24_000_000_000, "revenue_currency": "VND", "nace_code": "C.13.20",
}


def _ready(client) -> str:
    client.post("/suppliers", json=PROFILE, headers=_auth())
    pid = client.get("/suppliers/sup-1/periods", headers=_auth()).json()["periods"][0]["id"]
    client.repo.add_activity_lines("sup-1", pid, [
        ActivityLine(id="a1", supplier_id="sup-1", source_doc_id="doc-1", site_id="main", energy_type=EnergyType.ELECTRICITY,
                     quantity=600_000, unit=Unit.KWH, period_start=date(2025, 1, 1), period_end=date(2025, 12, 31)),
        ActivityLine(id="a2", supplier_id="sup-1", source_doc_id="doc-2", site_id="main", energy_type=EnergyType.LPG,
                     quantity=10_000, unit=Unit.KG, period_start=date(2025, 1, 1), period_end=date(2025, 12, 31)),
    ])
    assert client.post(f"/suppliers/sup-1/periods/{pid}/compute", headers=_auth()).status_code == 200
    return pid


def test_vsme_needs_an_inventory(client):
    pid = _setup(client)
    assert client.get(f"/suppliers/sup-1/periods/{pid}/vsme", headers=_auth()).status_code == 409


def test_vsme_route_returns_the_filled_official_file_and_what_is_left(client):
    pid = _ready(client)
    r = client.get(f"/suppliers/sup-1/periods/{pid}/vsme", headers=_auth())
    assert r.status_code == 200
    body = r.json()
    cells = {f["cell"]: f["value"] for f in body["filled"]}
    assert cells["'General Information'!D3"] == "Hoa Sen Garments"
    assert cells["'General Information'!E287"] == "Viet Nam"
    assert cells["'Environmental Disclosures'!D22"] == 29.4          # 10 t of LPG x 2.94 kg per kg
    assert {o["ref"] for o in body["open_items"]} >= {"B1 4", "B1 basis"}
    f = client.get(body["download"], headers=_auth())
    assert f.status_code == 200 and f.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert zipfile.ZipFile(io.BytesIO(f.content)).testzip() is None
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = load_workbook(io.BytesIO(f.content))
    assert wb["General Information"]["E284"].value == 420
    assert client.get(f"/suppliers/sup-1/periods/{pid}/vsme", headers=_auth("sup-2")).status_code == 403


def _request(client, **kw) -> str:
    return client.post("/suppliers/sup-1/requests", json={"customer_name": "Toyota Boshoku", "customer_share": 0.2, **kw}, headers=_auth()).json()["request"]["id"]


def test_pack_list_and_routes_need_the_right_inputs(client):
    _setup(client)
    rid = _request(client)
    names = [p["id"] for p in client.get("/suppliers/sup-1/packs", headers=_auth()).json()["packs"]]
    assert names == ["cdp_sme", "higg_fem_energy", "ecovadis_carbon"]
    assert client.get(f"/suppliers/sup-1/requests/{rid}/packs/cdp_sme", headers=_auth()).status_code == 409      # no inventory yet
    assert client.get(f"/suppliers/sup-1/requests/{rid}/packs/vsme_b3", headers=_auth()).status_code == 404       # an Excel bank, not a pack
    assert client.get("/suppliers/sup-1/requests/nope/packs/cdp_sme", headers=_auth()).status_code == 404


def test_pack_routes_return_the_pack_and_both_exports(client):
    pid = _ready(client)
    rid = _request(client, period_id=pid)
    r = client.get(f"/suppliers/sup-1/requests/{rid}/packs/cdp_sme", headers=_auth())
    assert r.status_code == 200
    body = r.json()
    items = {i["ref"]: i for i in body["pack"]["items"]}
    assert items["20.4"]["cells"][0]["value"] == 29.4 and items["20.12"]["cells"][0]["value"] == "Toyota Boshoku"
    assert sum(body["counts"].values()) == len(items)
    x = client.get(body["downloads"]["xlsx"], headers=_auth())
    assert x.status_code == 200 and x.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "attachment" in x.headers["content-disposition"]
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        assert load_workbook(io.BytesIO(x.content)).sheetnames[0] == "Answers"
    p = client.get(body["downloads"]["pdf"], headers=_auth())
    assert p.status_code == 200 and p.content[:5] == b"%PDF-"
    assert client.get(f"/suppliers/sup-1/requests/{rid}/packs/cdp_sme", headers=_auth("sup-2")).status_code == 403

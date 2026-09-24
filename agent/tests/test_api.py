"""API contract tests against the in-memory repo with auth disabled.

Model calls are avoided (no upload/reconcile here) so this runs offline; the
Gemini-backed paths are covered by the scripted demo and the golden eval.
"""

import io
from datetime import date

import pytest
from fastapi.testclient import TestClient

from app.models import ActivityLine, EnergyType, Unit
from app.repo import MemoryRepo, set_repo


@pytest.fixture
def client(monkeypatch):
    from app import settings as s

    s.get_settings.cache_clear()
    monkeypatch.setenv("AUTH_DISABLED", "TRUE")
    monkeypatch.setenv("STORAGE_BACKEND", "memory")
    from app import main

    repo = MemoryRepo()
    set_repo(repo)
    from app import storage as st
    st._storage = None
    c = TestClient(main.app)
    c.repo = repo
    yield c
    s.get_settings.cache_clear()


def _auth(uid="sup-1"):
    return {"Authorization": f"Bearer {uid}"}


PROFILE = {"fiscal_year_start": "2025-01-01", "fiscal_year_end": "2025-12-31", "country": "VN"}


def _setup(client) -> str:
    """Create the supplier and return the id of the default period it got."""
    client.post("/suppliers", json=PROFILE, headers=_auth())
    periods = client.get("/suppliers/sup-1/periods", headers=_auth()).json()["periods"]
    assert len(periods) == 1 and periods[0]["label"] == "FY2025"
    return periods[0]["id"]


def _lines(pid: str):
    return [
        ActivityLine(id="a1", supplier_id="sup-1", source_doc_id="doc-1", site_id="main", energy_type=EnergyType.ELECTRICITY,
                     quantity=10_000, unit=Unit.KWH, period_start=date(2025, 1, 1), period_end=date(2025, 1, 31)),
        ActivityLine(id="a2", supplier_id="sup-1", source_doc_id="doc-2", site_id="main", energy_type=EnergyType.DIESEL,
                     quantity=1000, unit=Unit.LITRE, period_start=date(2025, 2, 1), period_end=date(2025, 2, 1)),
    ]


def test_health(client):
    r = client.get("/health")
    assert r.status_code == 200 and r.json()["ok"] is True


def test_supplier_upsert_creates_default_period_and_keeps_site_ids(client):
    body = {**PROFILE, "name": "Minh Phát", "language": "vi", "sites": [{"name": "Nhơn Trạch", "country": "VN"}]}
    r = client.post("/suppliers", json=body, headers=_auth())
    assert r.status_code == 200
    p = r.json()
    assert p["id"] == "sup-1" and p["country"] == "VN" and len(p["sites"]) == 1
    site_id = p["sites"][0]["id"]
    # Re-saving with the same site name keeps its id, so earlier bills stay attached.
    r = client.post("/suppliers", json={**body, "sites": [{"name": "nhơn trạch", "country": "VN", "region": "SE"}]}, headers=_auth())
    assert r.json()["sites"][0]["id"] == site_id and r.json()["sites"][0]["region"] == "SE"
    # One default period, not one per save.
    assert len(client.get("/suppliers/sup-1/periods", headers=_auth()).json()["periods"]) == 1
    assert client.get("/suppliers/sup-1", headers=_auth("sup-2")).status_code == 403


def test_periods_dedupe_and_validate(client):
    pid = _setup(client)
    same = client.post("/suppliers/sup-1/periods", json={"start": "2025-01-01", "end": "2025-12-31"}, headers=_auth()).json()
    assert same["id"] == pid
    other = client.post("/suppliers/sup-1/periods", json={"start": "2024-04-01", "end": "2025-03-31"}, headers=_auth()).json()
    assert other["label"] == "FY2024/25" and other["id"] != pid
    assert client.post("/suppliers/sup-1/periods", json={"start": "2025-01-01", "end": "2024-12-31"}, headers=_auth()).status_code == 400
    assert client.get("/suppliers/sup-1/periods/nope", headers=_auth()).status_code == 404


def test_compute_and_inventory_are_period_scoped(client):
    pid = _setup(client)
    client.repo.add_activity_lines("sup-1", pid, _lines(pid))
    r = client.post(f"/suppliers/sup-1/periods/{pid}/compute", headers=_auth())
    assert r.status_code == 200
    t = r.json()["totals"]
    assert t["scope1_kg"] == pytest.approx(2660) and t["scope2_location_kg"] == pytest.approx(6592, rel=1e-3)

    inv = client.get(f"/suppliers/sup-1/periods/{pid}/inventory", headers=_auth()).json()
    assert len(inv["lines"]) == 3
    for l in inv["lines"]:
        assert l["source_doc_id"] and l["factor_id"] and l["factor_source_name"]
    assert inv["snapshot"]["period_id"] == pid
    assert set(inv["snapshot"]["totals"]["by_month"]) == {"2025-01", "2025-02"}

    # A second period on the same supplier sees none of it.
    other = client.post("/suppliers/sup-1/periods", json={"start": "2024-01-01", "end": "2024-12-31"}, headers=_auth()).json()
    inv2 = client.get(f"/suppliers/sup-1/periods/{other['id']}/inventory", headers=_auth()).json()
    assert inv2["snapshot"] is None and inv2["lines"] == []


def test_patch_line_marks_secondary_and_recomputes(client):
    pid = _setup(client)
    client.repo.add_activity_lines("sup-1", pid, _lines(pid)[1:])
    client.post(f"/suppliers/sup-1/periods/{pid}/compute", headers=_auth())
    r = client.patch(f"/suppliers/sup-1/periods/{pid}/activity-lines/a2", json={"quantity": 1200}, headers=_auth())
    assert r.status_code == 200
    assert r.json()["snapshot"]["totals"]["scope1_kg"] == pytest.approx(1200 * 2.66)
    line = client.repo.list_activity_lines("sup-1", pid)[0]
    assert line.tier == "secondary" and "corrected by supplier" in line.notes


def test_requests_derive_status_from_shared_period(client):
    pid = _setup(client)
    # Two customers asking for the same year share the period.
    a = client.post("/suppliers/sup-1/requests", json={"customer_name": "Aiko Corp", "due_date": "2026-01-31", "customer_share": 0.2}, headers=_auth()).json()
    b = client.post("/suppliers/sup-1/requests", json={"customer_name": "Toyota Boshoku", "period_id": pid}, headers=_auth()).json()
    assert a["period"]["id"] == pid == b["period"]["id"]
    assert a["status"] == "new" and a["shared_with"] == []  # created first, alone on the year
    assert b["shared_with"] == ["Aiko Corp"]
    assert client.get(f"/suppliers/sup-1/requests/{a['request']['id']}", headers=_auth()).json()["shared_with"] == ["Toyota Boshoku"]
    assert client.post("/suppliers/sup-1/requests", json={"customer_name": "  "}, headers=_auth()).status_code == 400

    # Bills reach the period: both requests move to "bills".
    from app.repo import DocumentRecord
    client.repo.put_document(DocumentRecord(id="doc-1", supplier_id="sup-1", period_id=pid, storage_key="k", filename="a.png", mime_type="image/png",
                                            status="extracted", extraction={"lines": [{"energy_type": "electricity"}]}))
    client.repo.add_activity_lines("sup-1", pid, _lines(pid))
    assert client.get(f"/suppliers/sup-1/requests/{a['request']['id']}", headers=_auth()).json()["status"] == "bills"
    assert client.get(f"/suppliers/sup-1/requests/{b['request']['id']}", headers=_auth()).json()["status"] == "bills"

    # Inventory computed: "inventory" for both; a third request on another year stays "new".
    client.post(f"/suppliers/sup-1/periods/{pid}/compute", headers=_auth())
    c = client.post("/suppliers/sup-1/requests", json={"customer_name": "Old Buyer", "new_period": {"start": "2024-01-01", "end": "2024-12-31"}}, headers=_auth()).json()
    home = client.get("/suppliers/sup-1/home", headers=_auth()).json()
    statuses = {r["request"]["customer_name"]: r["status"] for r in home["requests"]}
    assert statuses == {"Aiko Corp": "inventory", "Toyota Boshoku": "inventory", "Old Buyer": "new"}
    assert len(home["periods"]) == 2
    fy25 = next(p for p in home["periods"] if p["period"]["id"] == pid)
    assert sorted(fy25["used_by"]) == ["Aiko Corp", "Toyota Boshoku"] and fy25["totals"]["scope1_kg"] == pytest.approx(2660)
    assert fy25["progress"] == {"bills": True, "questions": True, "inventory": True, "documents": 1, "reconcile_status": "idle"}

    # Patch and delete.
    r = client.patch(f"/suppliers/sup-1/requests/{a['request']['id']}", json={"customer_name": "Aiko Corporation", "clear_due_date": True}, headers=_auth()).json()
    assert r["request"]["customer_name"] == "Aiko Corporation" and r["request"]["due_date"] is None
    assert client.delete(f"/suppliers/sup-1/requests/{c['request']['id']}", headers=_auth()).status_code == 204
    assert client.get(f"/suppliers/sup-1/requests/{c['request']['id']}", headers=_auth()).status_code == 404
    # The period the deleted request created is still there for reuse.
    assert len(client.get("/suppliers/sup-1/periods", headers=_auth()).json()["periods"]) == 2
    assert client.get(f"/suppliers/sup-1/requests/{a['request']['id']}", headers=_auth("sup-2")).status_code == 403


def test_questionnaire_and_methodology_need_inventory_first(client):
    pid = _setup(client)
    req = client.post("/suppliers/sup-1/requests", json={"customer_name": "Aiko"}, headers=_auth()).json()["request"]
    r = client.post(f"/suppliers/sup-1/requests/{req['id']}/questionnaire", files={"file": ("q.xlsx", io.BytesIO(b"x"), "application/octet-stream")}, headers=_auth())
    assert r.status_code == 409
    assert client.get(f"/suppliers/sup-1/periods/{pid}/methodology", headers=_auth()).status_code == 409


def _seed_reconciled_period(client) -> str:
    """A period that has one bill, one answered gap question, and a computed inventory."""
    from app.models import Assumption
    from app.repo import DocumentRecord, ReconcileState

    pid = _setup(client)
    extraction = {
        "doc_type": "electricity_bill", "language": "vi", "vendor": "EVN", "customer": None, "site_hint": None,
        "currency": "VND", "total_amount": 1.0, "confidence": 1.0, "warnings": [],
        "lines": [{"energy_type": "electricity", "quantity": 10000, "unit": "kWh", "period_start": "2025-01-01",
                   "period_end": "2025-01-31", "gas": None, "confidence": 1.0, "evidence": "10,000 kWh"}],
    }
    client.repo.put_document(DocumentRecord(id="doc-1", supplier_id="sup-1", period_id=pid, storage_key=f"suppliers/sup-1/periods/{pid}/uploads/doc-1/a.png",
                                            filename="a.png", mime_type="image/png", site_id="main", status="extracted", extraction=extraction))
    client.repo.add_activity_lines("sup-1", pid, [
        ActivityLine(id="a1", supplier_id="sup-1", source_doc_id="doc-1", site_id="main", energy_type=EnergyType.ELECTRICITY,
                     quantity=12000, unit=Unit.KWH, period_start=date(2025, 1, 1), period_end=date(2025, 1, 31),
                     tier="secondary", notes="corrected by supplier from 10000 to 12000"),
        ActivityLine(id="a2", supplier_id="sup-1", source_doc_id="assumption:asm-1", site_id="main",
                     energy_type=EnergyType.ELECTRICITY, quantity=10000, unit=Unit.KWH, period_start=date(2025, 2, 1),
                     period_end=date(2025, 2, 28), tier="estimated", assumption_ids=["asm-1"]),
    ])
    client.repo.add_assumption("sup-1", pid, Assumption(id="asm-1", supplier_id="sup-1", question="q", answer="a", applied_to=["a2"], method="m"))
    client.repo.put_reconcile(ReconcileState(supplier_id="sup-1", period_id=pid, status="done"))
    client.post(f"/suppliers/sup-1/periods/{pid}/compute", headers=_auth())
    assert client.get(f"/suppliers/sup-1/periods/{pid}/inventory", headers=_auth()).json()["snapshot"] is not None
    return pid


def test_reset_chat_rebuilds_lines_from_extraction_and_keeps_documents(client):
    pid = _seed_reconciled_period(client)
    r = client.delete(f"/suppliers/sup-1/periods/{pid}/data?scope=chat", headers=_auth())
    assert r.status_code == 200 and r.json()["activity_lines"] == 1
    lines = client.repo.list_activity_lines("sup-1", pid)
    assert len(lines) == 1 and lines[0].quantity == 10000 and lines[0].tier == "primary"
    assert client.repo.list_assumptions("sup-1", pid) == []
    view = client.get(f"/suppliers/sup-1/periods/{pid}/reconcile", headers=_auth()).json()
    assert view["status"] == "idle" and view["history"] == []
    elec = next(e for e in view["summary"]["energy"] if e["energy_type"] == "electricity")
    assert elec["documents"] == 1 and elec["coverage"][0] == "primary" and elec["months_missing"] == 11
    inv = client.get(f"/suppliers/sup-1/periods/{pid}/inventory", headers=_auth()).json()
    assert inv["snapshot"] is None and inv["lines"] == []
    assert len(client.get(f"/suppliers/sup-1/periods/{pid}/documents", headers=_auth()).json()["documents"]) == 1


def test_delete_document_removes_lines_and_resets_inventory(client, tmp_path):
    from app import storage as st

    st._storage = st.LocalStorage(tmp_path)
    pid = _seed_reconciled_period(client)
    st._storage.put(f"suppliers/sup-1/periods/{pid}/uploads/doc-1/a.png", b"png")
    r = client.delete(f"/suppliers/sup-1/periods/{pid}/documents/doc-1", headers=_auth())
    assert r.status_code == 204
    assert client.get(f"/suppliers/sup-1/periods/{pid}/documents", headers=_auth()).json()["documents"] == []
    assert {l.source_doc_id for l in client.repo.list_activity_lines("sup-1", pid)} == {"assumption:asm-1"}
    assert not st._storage.exists(f"suppliers/sup-1/periods/{pid}/uploads/doc-1/a.png")
    assert client.get(f"/suppliers/sup-1/periods/{pid}/inventory", headers=_auth()).json()["snapshot"] is None
    assert client.get(f"/suppliers/sup-1/periods/{pid}/reconcile", headers=_auth()).json()["status"] == "idle"
    assert client.delete(f"/suppliers/sup-1/periods/{pid}/documents/doc-1", headers=_auth()).status_code == 404
    assert client.delete(f"/suppliers/sup-1/periods/{pid}/documents/doc-1", headers=_auth("sup-2")).status_code == 403


def test_reset_period_drops_shared_questionnaires_but_keeps_requests(client, tmp_path):
    from app import storage as st
    from app.repo import QuestionnaireRecord

    st._storage = st.LocalStorage(tmp_path)
    pid = _seed_reconciled_period(client)
    req = client.post("/suppliers/sup-1/requests", json={"customer_name": "Aiko"}, headers=_auth()).json()["request"]
    client.repo.put_questionnaire(QuestionnaireRecord(id="q1", supplier_id="sup-1", request_id=req["id"], period_id=pid, template_key="t", filled_key="f", status="filled"))
    client.repo.put_request("sup-1", client.repo.get_request("sup-1", req["id"]).model_copy(update={"questionnaire_id": "q1"}))
    assert client.get(f"/suppliers/sup-1/requests/{req['id']}", headers=_auth()).json()["status"] == "filled"

    assert client.delete(f"/suppliers/sup-1/periods/{pid}/data?scope=all", headers=_auth()).status_code == 200
    view = client.get(f"/suppliers/sup-1/requests/{req['id']}", headers=_auth()).json()
    assert view["status"] == "new" and view["request"]["questionnaire_id"] is None
    assert client.repo.get_questionnaire("sup-1", "q1") is None
    assert client.get(f"/suppliers/sup-1/periods/{pid}/documents", headers=_auth()).json()["documents"] == []
    assert client.delete(f"/suppliers/sup-1/periods/{pid}/data?scope=everything", headers=_auth()).status_code == 400

    # Whole-supplier reset keeps only the profile.
    assert client.delete("/suppliers/sup-1/data", headers=_auth()).status_code == 200
    assert client.get("/suppliers/sup-1/periods", headers=_auth()).json()["periods"] == []
    assert client.get("/suppliers/sup-1/requests", headers=_auth()).json()["requests"] == []
    assert client.get("/suppliers/sup-1", headers=_auth()).status_code == 200


def test_edit_answer_rewrites_workbook_and_locks_calculator_numbers(client, tmp_path):
    from openpyxl import Workbook, load_workbook

    from app.agents.filler import Answer
    from app.repo import QuestionnaireRecord
    from app import storage as st

    st._storage = st.LocalStorage(tmp_path)
    pid = _setup(client)

    wb = Workbook(); ws = wb.active; ws.title = "Q"
    ws.append(["Ref", "Question", "Response", "Notes"])
    ws.append(["20.4", "Scope 1 emissions (tCO2e)", None, None])
    ws.append(["20.16", "Do you have a reduction target?", None, None])
    buf = io.BytesIO(); wb.save(buf)
    st._storage.put("suppliers/sup-1/requests/r1/template.xlsx", buf.getvalue())

    answers = [
        Answer(question_id="q1", ref="20.4", question="Scope 1 emissions (tCO2e)", field="scope1_tco2e", value=12.5,
               source="calculator", sheet="Q", row=2, col=2, editable=False),
        Answer(question_id="q2", ref="20.16", question="Do you have a reduction target?", field="reduction_target",
               value=None, source="-", flagged=True, flag_reason="to be stated by the supplier", sheet="Q", row=3, col=2),
    ]
    rec = QuestionnaireRecord(id="q1", supplier_id="sup-1", request_id="r1", period_id=pid, template_key="suppliers/sup-1/requests/r1/template.xlsx",
                              filled_key="suppliers/sup-1/requests/r1/outputs/q1/filled.xlsx", answers=[a.model_dump() for a in answers], status="filled")
    client.repo.put_questionnaire(rec)

    r = client.patch("/suppliers/sup-1/questionnaire/q1/answers/q1", json={"value": "999"}, headers=_auth())
    assert r.status_code == 400 and "inventory" in r.json()["detail"]

    r = client.patch("/suppliers/sup-1/questionnaire/q1/answers/q2", json={"value": "Cut Scope 2 by 30% by 2030 via rooftop solar."}, headers=_auth())
    assert r.status_code == 200
    a = next(x for x in r.json()["questionnaire"]["answers"] if x["question_id"] == "q2")
    assert a["flagged"] is False and a["edited"] is True and a["source"] == "entered by supplier"

    out = load_workbook(io.BytesIO(st._storage.get("suppliers/sup-1/requests/r1/outputs/q1/filled.xlsx")))["Q"]
    assert out["C3"].value.startswith("Cut Scope 2") and out["C3"].fill.fgColor.rgb.endswith("DCEBFA")
    assert out["C2"].value == 12.5

    r = client.patch("/suppliers/sup-1/questionnaire/q1/answers/q2", json={"value": "Net zero by 2040."}, headers=_auth())
    assert "generated answer was" not in r.json()["questionnaire"]["answers"][1]["source"]
    assert client.patch("/suppliers/sup-1/questionnaire/q1/answers/q2", json={"value": "  "}, headers=_auth()).status_code == 400
    assert client.patch("/suppliers/sup-1/questionnaire/q1/answers/zz", json={"value": "x"}, headers=_auth()).status_code == 404
    assert client.patch("/suppliers/sup-1/questionnaire/q1/answers/q2", json={"value": "x"}, headers=_auth("sup-2")).status_code == 403


def test_unauthenticated_is_rejected_when_auth_enabled(monkeypatch):
    from app import settings as s
    s.get_settings.cache_clear()
    monkeypatch.setenv("AUTH_DISABLED", "FALSE")
    from app import main
    c = TestClient(main.app)
    assert c.get("/suppliers/x").status_code == 401
    s.get_settings.cache_clear()
